"""
Printing station views for IntelliScale.
Handles local printing station operations, printing notes, and records.
"""
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.core.paginator import Paginator
from django.utils import timezone
from ..models import Scale, Product, PrintingNote, PrintingRecord


@login_required
def printing_station(request):
    """
    View for the local printing station.
    - Handles scale interaction (via existing endpoints)
    - Manages PrintingNotes and PrintingRecords
    - purely local, no ERP/Odoo interaction
    """
    if request.method == 'POST':
        # Handle form submission
        printing_note_id = request.POST.get('printing_note_id')
        grower_number = request.POST.get('grower_number')
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        total_bales = request.POST.get('expected_bales')
        
        # Scale/Record Data
        scale_id = request.POST.get('scale_id')
        product_id = request.POST.get('product_id')
        barcode = request.POST.get('barcode', '')
        gross_weight = request.POST.get('gross_weight') or 0
        tare_weight = request.POST.get('tare_weight') or 0
        net_weight = request.POST.get('net_weight') or 0
        
        moisture = request.POST.get('moisture') or None
        
        # Validate moisture is a valid percentage (0-100)
        if moisture is not None:
            try:
                moisture_float = float(moisture)
                if moisture_float < 0 or moisture_float > 100:
                    messages.error(request, "Moisture must be between 0 and 100%.")
                    return redirect('scale:printing_station')
            except (ValueError, TypeError):
                messages.error(request, "Invalid moisture value.")
                return redirect('scale:printing_station')
        
        # Find or Create Note
        printing_note = None
        if printing_note_id:
            try:
                printing_note = PrintingNote.objects.get(id=printing_note_id, user=request.user)
            except PrintingNote.DoesNotExist:
                pass # Should not happen unless tampering or new session
        
        if not printing_note:
            # Create new note
            printing_note = PrintingNote.objects.create(
                user=request.user,
                grower_number=grower_number,
                first_name=first_name,
                last_name=last_name,
                expected_bales=total_bales if total_bales else None
            )
        else:
            printing_note.grower_number = grower_number or ''
            printing_note.first_name = first_name or ''
            printing_note.last_name = last_name or ''
            if total_bales is not None: printing_note.expected_bales = total_bales if total_bales else None
            printing_note.save() # Updates updated_at
            
        # Update Session with Active Note ID
        request.session['active_printing_note_id'] = printing_note.id
            
        # Create Record
        product = None
        if product_id:
            try:
                product = Product.objects.get(id=product_id)
            except Product.DoesNotExist:
                pass

        # Look up Scale object to get its scale_id
        scale = None
        if scale_id:
            try:
                scale = Scale.objects.get(id=scale_id)
            except (Scale.DoesNotExist, ValueError):
                pass
                
        if not barcode:
            messages.error(request, "Barcode cannot be empty.")
            return redirect('scale:printing_station')

        try:
            net_weight_float = float(net_weight)
        except (ValueError, TypeError):
            net_weight_float = 0

        if net_weight_float <= 0:
            messages.error(request, "Net weight must be greater than zero.")
            return redirect('scale:printing_station')

        # Check if barcode already exists in this note
        existing_record = PrintingRecord.objects.filter(printing_note=printing_note, barcode=barcode).first()
        rescan = request.POST.get('rescan')

        if existing_record:
            if rescan == 'true':
                # Update existing record
                existing_record.scale_id = scale.scale_id if scale else scale_id
                existing_record.product = product
                existing_record.gross_weight = gross_weight
                existing_record.tare_weight = tare_weight
                existing_record.net_weight = net_weight
                existing_record.moisture = moisture
                # Update timestamp to bring it to top of list
                existing_record.timestamp = timezone.now()
                existing_record.save()
                
                messages.success(request, f"Updated {barcode} ({net_weight} kg)")
                return redirect('scale:printing_station')
            else:
                messages.error(request, f"Barcode {barcode} already exists in this note.")
                return redirect('scale:printing_station')



        PrintingRecord.objects.create(
            printing_note=printing_note,
            scale_id=scale.scale_id if scale else scale_id,
            product=product,
            barcode=barcode,
            gross_weight=gross_weight,
            tare_weight=tare_weight,
            net_weight=net_weight,
            moisture=moisture,
            unit_of_measure='kg' # Default for now
        )
        
        messages.success(request, f"Recorded {barcode} ({net_weight} kg)")
        return redirect('scale:printing_station')

    else:
        # GET Request
        new_session = request.GET.get('new_session')
        printing_note = None
        
        if new_session == 'true':
            # Intentional new session, clear active note in session
            if 'active_printing_note_id' in request.session:
                del request.session['active_printing_note_id']
        else:
            # Try to resume from session
            active_note_id = request.session.get('active_printing_note_id')
            if active_note_id:
                try:
                    printing_note = PrintingNote.objects.get(id=active_note_id, user=request.user)
                except PrintingNote.DoesNotExist:
                    # Session ID invalid (maybe deleted), clear it
                    del request.session['active_printing_note_id']
            
        # Create a dictionary of product tare weights for the template
        product_tare_weights = {}
        processed_products = Product.objects.filter(is_active=True)
        for product in processed_products:
            product_tare_weights[product.id] = float(product.tare_weight or 0)

        context = {
            'scales': Scale.objects.filter(is_active=True),
            'products': processed_products,
            'printing_note': printing_note,
            'product_tare_weights': product_tare_weights,
            'records': PrintingRecord.objects.filter(printing_note=printing_note).order_by('-timestamp'),
            # Pass today's date for display if needed
        }
        return render(request, 'scale/printing_station.html', context)


@login_required
def printing_note_list(request):
    """
    List all printing notes.
    """
    notes_list = PrintingNote.objects.all().order_by('-created_at')
    
    # Filtering
    grower_number = request.GET.get('grower_number')
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    
    if grower_number:
        notes_list = notes_list.filter(grower_number__icontains=grower_number)
    
    if start_date:
        notes_list = notes_list.filter(created_at__date__gte=start_date)
        
    if end_date:
        notes_list = notes_list.filter(created_at__date__lte=end_date)

    paginator = Paginator(notes_list, 20) # Show 20 notes per page
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    return render(request, 'scale/printing_note_list.html', {'page_obj': page_obj})


@login_required
def printing_note_detail(request, pk):
    """
    Detail view for a printing note.
    """
    note = get_object_or_404(PrintingNote, pk=pk)
    return render(request, 'scale/printing_note_detail.html', {'note': note})


@login_required
def printing_note_delete(request, pk):
    """
    Delete a printing note.
    """
    note = get_object_or_404(PrintingNote, pk=pk)
    
    if request.method == 'POST':
        note.delete()
        messages.success(request, f'Printing note #{pk} deleted successfully.')
        return redirect('scale:printing_note_list')
    
    # If not POST, redirect back to list (though this shouldn't be reached if only used via button)
    return redirect('scale:printing_note_list')


@login_required
def reactivate_printing_note(request, pk):
    """
    Reactivate a printing note for editing in the printing station.
    """
    note = get_object_or_404(PrintingNote, pk=pk)
    
    # Set this note as the active one in the session
    request.session['active_printing_note_id'] = note.id
    
    messages.success(request, f'Printing Note #{note.id} reactivated.')
    return redirect('scale:printing_station')


@login_required
def update_printing_note_details(request):
    """
    AJAX endpoint to save printing note details without requiring a barcode scan.
    """
    if request.method == 'POST':
        note_id = request.POST.get('printing_note_id')
        if not note_id:
            return JsonResponse({'success': False, 'message': 'No note ID provided'})

        try:
            note = PrintingNote.objects.get(pk=note_id, user=request.user)
        except PrintingNote.DoesNotExist:
            return JsonResponse({'success': False, 'message': 'Note not found'})

        note.grower_number = request.POST.get('grower_number', '')
        note.first_name = request.POST.get('first_name', '')
        note.last_name = request.POST.get('last_name', '')
        expected_bales = request.POST.get('expected_bales')
        note.expected_bales = expected_bales if expected_bales else None
        note.save()
        return JsonResponse({'success': True, 'message': 'Details saved successfully'})

    return JsonResponse({'success': False, 'message': 'Invalid request method'})


@login_required
def printing_record_delete(request, pk):
    record = get_object_or_404(PrintingRecord, pk=pk)
    note_id = record.printing_note.id
    
    if request.method == 'POST':
        barcode = record.barcode
        record.delete()
        messages.success(request, f'Record for barcode {barcode} deleted successfully.')
    
    return redirect('scale:printing_note_detail', pk=note_id)


@login_required
def printing_note_export_xlsx(request, pk):
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    note = get_object_or_404(PrintingNote, pk=pk)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"Printing_Note_{note.id}_{note.created_at.strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Note {note.id}"

    # Styles
    header_font = Font(bold=True, size=12)
    title_font = Font(bold=True, size=14)
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    thin_border = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))

    # Note Information
    ws['A1'] = "Printing Note Details"
    ws['A1'].font = title_font
    ws.merge_cells('A1:F1')
    ws['A1'].alignment = center_align

    # Metadata rows
    # Round totals to 2 decimal places
    total_gross = float(note.get_total_gross_weight() or 0)
    total_net = float(note.get_total_net_weight() or 0)
    
    metadata = [
        ("Note ID:", str(note.id)),
        ("Grower Number:", note.grower_number),
        ("Grower Name:", f"{note.first_name} {note.last_name}"),
        ("Created At:", note.created_at.strftime('%Y-%m-%d %H:%M')),
        ("Expected Bales:", str(note.expected_bales or "-")),
        ("Received Bales:", str(note.records.count())),
        ("Total Gross:", f"{total_gross:.2f} kg"),
        ("Total Net:", f"{total_net:.2f} kg"),
    ]

    row_num = 3
    for label, value in metadata:
        ws.cell(row=row_num, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row_num, column=2, value=value)
        row_num += 1

    row_num += 2  # Gap

    # Table Headers
    headers = ['Barcode', 'Product', 'Scale', 'Gross', 'Tare', 'Net', 'Moisture', 'Timestamp']
    ws.append([]) # Empty row space if needed or just set start row
    
    # Reset row_num for table
    table_start_row = row_num
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=table_start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        # simple gray fill
        cell.fill = PatternFill(start_color="EEEEEE", end_color="EEEEEE", fill_type="solid")

    # Table Data
    row_num = table_start_row + 1
    for record in note.records.all():
        data = [
            record.barcode,
            record.product.name if record.product else "-",
            record.scale_id,
            f"{record.gross_weight} {record.unit_of_measure}",
            f"{record.tare_weight} {record.unit_of_measure}",
            f"{record.net_weight} {record.unit_of_measure}",
            f"{record.moisture}%" if record.moisture else "0%",
            record.timestamp.strftime('%Y-%m-%d %H:%M')
        ]
        
        for col_idx, value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.alignment = left_align
            cell.border = thin_border
            if col_idx in [4, 5, 6, 7]: # Weight columns and Moisture
                cell.alignment = right_align
        
        row_num += 1

    # Adjust column widths
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    wb.save(response)
    return response
